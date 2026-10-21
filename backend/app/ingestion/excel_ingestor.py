import time
from pathlib import Path
from typing import List, Tuple
import openpyxl
from app.ingestion.models import DocumentMetadata, DocumentChunk, IngestionReport

class ExcelIngestor:
    """Parses Excel spreadsheets (.xlsx, .xlsm) using openpyxl, preserving sheets, headers, and cell provenance."""

    def __init__(self, rows_per_chunk: int = 5):
        self.rows_per_chunk = rows_per_chunk

    def parse(self, file_path: Path, metadata: DocumentMetadata) -> Tuple[List[DocumentChunk], IngestionReport]:
        start_time = time.time()
        chunks: List[DocumentChunk] = []
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)

        sheet_count = len(wb.sheetnames)
        metadata.page_count = sheet_count

        for sheet_idx, sheet_name in enumerate(wb.sheetnames, start=1):
            sheet = wb[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue

            # Identify headers (first row with >= 2 non-empty values)
            header_idx = 0
            headers = []
            for idx, r in enumerate(rows):
                non_empty = [str(c).strip() for c in r if c is not None and str(c).strip() != ""]
                if len(non_empty) >= 2:
                    header_idx = idx
                    headers = [str(c).strip() if c is not None else f"Column_{i+1}" for i, c in enumerate(r)]
                    break

            data_rows = rows[header_idx + 1:]
            if not data_rows and headers:
                # Only header exists
                chunk_id = f"{metadata.doc_id}_s{sheet_idx}_h0"
                chunk_text = f"Spreadsheet [{metadata.filename}], Sheet [{sheet_name}]\nColumns: {', '.join(headers)}"
                chunks.append(DocumentChunk(
                    chunk_id=chunk_id,
                    doc_id=metadata.doc_id,
                    text=chunk_text,
                    sheet_name=sheet_name,
                    source_type="excel_sheet",
                    cell_range=f"Row 1 to {header_idx + 1}",
                    metadata={"headers": headers, "sheet_index": sheet_idx}
                ))
                continue

            # Group rows into structured chunks
            for i in range(0, len(data_rows), self.rows_per_chunk):
                batch = data_rows[i:i + self.rows_per_chunk]
                start_row_num = header_idx + 2 + i
                end_row_num = start_row_num + len(batch) - 1

                lines = [f"Spreadsheet Record: [{metadata.filename}] | Sheet: [{sheet_name}] | Rows: {start_row_num}-{end_row_num}"]
                for row_offset, row in enumerate(batch):
                    actual_row_num = start_row_num + row_offset
                    row_parts = []
                    for col_idx, cell_value in enumerate(row):
                        if cell_value is not None and str(cell_value).strip() != "":
                            col_name = headers[col_idx] if col_idx < len(headers) else f"Col_{col_idx+1}"
                            row_parts.append(f"{col_name}: {cell_value}")
                    if row_parts:
                        lines.append(f"- Row {actual_row_num}: " + " | ".join(row_parts))

                chunk_text = "\n".join(lines)
                chunk_id = f"{metadata.doc_id}_s{sheet_idx}_r{start_row_num}_{end_row_num}"
                chunks.append(DocumentChunk(
                    chunk_id=chunk_id,
                    doc_id=metadata.doc_id,
                    text=chunk_text,
                    sheet_name=sheet_name,
                    source_type="excel_sheet",
                    cell_range=f"Rows {start_row_num}-{end_row_num}",
                    token_count=len(chunk_text.split()),
                    metadata={
                        "headers": headers,
                        "sheet_name": sheet_name,
                        "start_row": start_row_num,
                        "end_row": end_row_num,
                        "filename": metadata.filename
                    }
                ))

        wb.close()
        duration_ms = (time.time() - start_time) * 1000.0
        report = IngestionReport(
            doc_id=metadata.doc_id,
            filename=metadata.filename,
            file_type="excel",
            chunks_created=len(chunks),
            pages_or_sheets=sheet_count,
            parser_used="openpyxl",
            processing_time_ms=duration_ms
        )
        return chunks, report
